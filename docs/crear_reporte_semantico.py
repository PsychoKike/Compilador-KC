from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from sintac import parse_code
from semantico import analizar_semantica

doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(.75)
s.left_margin = s.right_margin = Inches(.8)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(29)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 2'].font.size = Pt(12)
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
header = s.header.paragraphs[0]
header.text = 'COMPILADOR KC   |   COMPILADORES II   |   FASE 3'
header.runs[0].font.size = Pt(9)
footer = s.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Análisis semántico  ·  ')
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

def p(text, style=None):
    return doc.add_paragraph(text, style)

def h(text, level=1):
    doc.add_heading(text, level)

def page(title):
    doc.add_page_break(); h(title)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, text in enumerate(headers): t.rows[0].cells[i].text = str(text)
    for row in rows:
        for c, text in zip(t.add_row().cells, row): c.text = str(text)
    for ri, row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); pr.append(no_split)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); pr.append(repeat)
        for ci, c in enumerate(row.cells):
            if widths: c.width = Inches(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                b = OxmlElement('w:' + edge)
                for k,v in {'val':'single','sz':'4','color':'D9D9D9'}.items(): b.set(qn('w:'+k),v)
                borders.append(b)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                m = OxmlElement('w:'+edge); m.set(qn('w:w'),'85'); m.set(qn('w:type'),'dxa'); margins.append(m)
            tcpr.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7EDF3' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF')); tcpr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing = 1
                for r in para.runs: r.font.size = Pt(9); r.bold = ri==0
    p('')
    return t

def code(text):
    para = p(text)
    para.paragraph_format.space_after = Pt(9)
    para.paragraph_format.line_spacing = 1
    for r in para.runs: r.font.name='Consolas'; r.font.size=Pt(9)

def capture(number, title, guidance):
    h(f'Evidencia {number} {title}', 2)
    p(guidance)
    para = p(f'[Insertar aquí la captura {number}]')
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(42)
    para.paragraph_format.space_after = Pt(55)
    p(f'Figura {number}. {title}. Fuente: ejecución del Compilador KC por el equipo.')

p('Universidad Autónoma de Aguascalientes')
p('Ingeniería en Sistemas Computacionales\nMateria Compiladores II')
p('Documentación técnica del análisis semántico', 'Title')
p('Compilador KC\nFase 3 del proyecto', 'Subtitle')
p('Docente Dra. Blanca Guadalupe Estrada Rentería')
p('Integrantes y matrículas\n[Nombre completo]  [Matrícula]\n[Nombre completo]  [Matrícula]\n[Nombre completo]  [Matrícula]')
p('Grupo [Completar]     Fecha de entrega [Completar]')
h('Propósito del documento', 2)
p('Describimos cómo incorporamos el análisis semántico al Compilador KC, cómo consultamos la tabla de símbolos y cómo propagamos los tipos y valores sobre una copia del árbol sintáctico. Documentamos las reglas implementadas y los resultados reproducidos con el programa de prueba proporcionado.')
p('Esta versión de Word es editable. Los espacios de evidencia quedan pendientes de completar con capturas del IDE antes de exportar el documento a PDF para la entrega.')

page('1 Objetivo y organización de la implementación')
p('El análisis semántico comprueba que las construcciones aceptadas por el parser tengan significado según las reglas del lenguaje. Reconocer una asignación no basta: también se necesita verificar que el identificador exista y que el tipo de la expresión sea compatible con la variable destino.')
p('Implementamos esta etapa en Python y la integramos al IDE existente de Tkinter. Reutilizamos el analizador léxico y el parser basado en PLY. La entrada semántica es un AST; la salida contiene el árbol anotado, la tabla de símbolos y una lista de errores con código, línea, columna y descripción.')
table(['Archivo', 'Responsabilidad'], [
('lexer.py','Produce tokens y conserva la ubicación del código fuente.'),
('sintac.py','Construye ASTNode y entrega el AST mediante parse_code.'),
('semantico.py','Define Simbolo, TablaSimbolos y AnalizadorSemantico.'),
('file_operations.py','Coordina los análisis y presenta árbol, tabla y errores.'),
('editor.py y buttons.py','Conectan los botones y las pestañas del IDE.'),
('test_semantico.py','Verifica los comportamientos semánticos con unittest.'),
('docs/diseno_semantico.md','Describe las reglas de tipos y los atributos del diseño.')], [1.9,5])
h('Flujo general',2)
code('Código fuente\n  → tokens del lexer\n  → AST del parser\n  → copia profunda del AST\n  → recorrido semántico y consultas a la tabla\n  → árbol anotado + tabla de símbolos + errores')
p('Los botones Run y Semántico ejecutan las fases en secuencia. El segundo también selecciona las pestañas del resultado. La existencia de errores no impide mostrar el árbol disponible ni continuar revisando otras sentencias.')

page('2 Estructura del árbol y atributos semánticos')
p('Cada ASTNode contiene type, children, leaf, level, lineno y col. type identifica la construcción u operador; children mantiene sus descendientes; leaf conserva el identificador o literal; lineno y col permiten ubicar el diagnóstico. level se utiliza para representar la jerarquía.')
p('En analizar utilizamos copy.deepcopy(ast). Las anotaciones se agregan a la copia y el AST original permanece disponible para compararlo con la pestaña Semántico. No sustituimos el árbol por una representación independiente.')
table(['Atributo', 'Origen y uso'],[
('dtype','Heredado. El tipo de la declaración se transmite a cada identificador.'),
('tipo','Sintetizado. Proviene de literales o de la tabla y sube por las expresiones.'),
('valor','Sintetizado. Se calcula desde los operandos; None representa valor desconocido.'),
('tipo_esperado','Heredado. La asignación establece su tipo destino y una condición espera bool.'),
('conversion','Marca int→float en operandos mixtos o en la asignación.'),
('ambito y direccion','Relacionan identificadores con su contexto y desplazamiento.'),
('error_sem','Guarda el código del diagnóstico asociado al nodo.'),
('actualizacion y postfijo','El parser conserva el cambio ±1 y si el operador es postfijo.')],[1.6,5.3])
h('Atributos heredados y sintetizados',2)
p('En int x, y, z, el tipo int se propaga desde el nodo de tipo hacia las tres hojas mediante dtype. En 2 + 3 - 1, los valores de las hojas se propagan hacia la suma y después hacia la resta. La asignación recibe el resultado y actualiza la entrada del símbolo destino.')
p('El atributo tipo_esperado registra el contexto requerido. La comprobación efectiva de la compatibilidad se realiza en los métodos de asignación y validación de condiciones; la anotación por sí sola no decide la validez.')

page('3 Construcción y uso de la tabla de símbolos')
p('TablaSimbolos mantiene una pila de ámbitos. Cada elemento contiene el nombre del ámbito y un diccionario cuya clave es el lexema y cuyo valor es un objeto Simbolo. Esta estructura permite usar las búsquedas de diccionario de Python y resolver primero los nombres más cercanos.')
table(['Campo', 'Información almacenada'],[
('nombre y tipo','Lexema y tipo normalizado int, float o bool; real se normaliza a float.'),
('lineas y col','Primera línea de declaración, líneas de uso sin repetir y columna de declaración.'),
('direccion','Desplazamiento estático asignado por siguiente_dir.'),
('ambito','Nombre del contexto en el que fue declarado el símbolo.'),
('valor','Último valor conocido durante el análisis; inicia en None.'),
('revision','Contador de escrituras usado para el seguimiento conservador de ciclos.')],[1.6,5.3])
table(['Operación', 'Cómo interviene durante el análisis'],[
('entrar_ambito y salir_ambito','Abren y cierran contextos para main, bloques, ramas y ciclos.'),
('buscar_local','Consulta únicamente el contexto actual para detectar duplicados.'),
('buscar','Recorre la pila desde el ámbito interno hasta el global.'),
('insertar','Crea el símbolo, reserva su dirección y lo conserva en entradas.'),
('registrar_uso','Agrega una línea si todavía no aparece en el símbolo.')],[2,4.9])
p('Las variables int y float reservan 4 bytes y bool reserva 1. siguiente_dir comienza en cero y aumenta después de cada inserción. Son desplazamientos del modelo del compilador, no direcciones reales de memoria de Python.')
p('La lista entradas conserva también los símbolos de ámbitos cerrados para mostrarlos en el IDE. En cambio, la búsqueda semántica utiliza solo la pila activa. Declarar dos veces en un mismo ámbito produce SEM02; un nombre local puede ocultar a otro de un ámbito exterior.')

page('4 Recorrido del árbol y procesamiento de declaraciones')
h('Punto de entrada y despacho',2)
p('analizar_semantica(ast) crea una instancia de AnalizadorSemantico y llama a analizar. Este método copia el AST, abre el ámbito main, visita la raíz, cierra el ámbito y devuelve los tres resultados. visitar identifica el tipo de sentencia y delega en el método correspondiente.')
code('analizar(ast):\n    anotado = copy.deepcopy(ast)\n    tabla.entrar_ambito("main")\n    visitar(anotado)\n    tabla.salir_ambito()\n    return anotado, tabla, errores')
p('Los contenedores main y lista_contenido se recorren de izquierda a derecha. Las sentencias mantienen el orden del programa, por lo que una asignación válida puede aportar un valor a expresiones posteriores. Los nodos de recuperación sintáctica se omiten como sentencias y se tratan como error cuando participan en una expresión.')
h('Declaraciones y usos de identificadores',2)
p('visitar_declaracion lee el tipo, lo normaliza mediante TIPOS y lo hereda a cada identificador. insertar verifica si el nombre ya existe en el ámbito actual. Si está disponible, crea el símbolo con valor None y asigna su desplazamiento; si ya existe, se registra SEM02 sin insertar otro símbolo.')
p('usar_identificador consulta buscar, registra la línea de uso y copia tipo, valor y ámbito al nodo. Si no encuentra el nombre, reporta SEM01 y establece el tipo interno error. Esta búsqueda se utiliza en destinos de asignación, operandos, condiciones, entrada, salida e incrementos.')
h('Recorrido de expresiones en postorden',2)
code('expresion(nodo):\n    si es identificador: consultar tabla\n    si es literal: obtener tipo y valor\n    en otro caso:\n        para cada hijo: expresion(hijo)\n        evaluar_operador(nodo)')
p('El operador se evalúa después de sus hijos. De esta manera recibe tipos y valores ya calculados. Los grupos entre paréntesis copian el tipo y valor de su expresión hija. No se utiliza eval sobre el texto fuente.')

page('5 Evaluación de expresiones y reglas de operadores')
p('Para evaluar usamos funciones del módulo operator, por ejemplo operator.add, operator.sub, operator.mul y operator.eq. Los diccionarios de operadores asocian el símbolo del lenguaje con la función correspondiente. Primero se comprueban los tipos y después se intenta calcular el valor.')
table(['Construcción', 'Regla implementada'],[
('int con int en aritmética','Resultado int. La división entera trunca hacia cero.'),
('Operación numérica mixta','Resultado float y marca int→float en el operando entero.'),
('Módulo %','Requiere dos enteros; de otro modo SEM05.'),
('Aritmética con bool','No es válida; genera SEM05.'),
('< <= > >=','Requieren operandos numéricos y producen bool.'),
('== y !=','Aceptan dos numéricos o dos bool; una mezcla bool y número produce SEM06.'),
('&& || and or not','Requieren operandos bool y producen bool.'),
('Potencia ^','Se evalúa con operator.pow y se ajusta al tipo numérico inferido.')],[2.1,4.8])
h('Ejemplo de evaluación paso a paso',2)
code('y = 2 + 3 - 1;\n\nHojas: 2 → int, valor 2; 3 → int, valor 3; 1 → int, valor 1\nNodo +: int + int → int, valor 5\nNodo -: int - int → int, valor 4\nNodo =: int compatible con int → y recibe 4')
p('Después, z = y + 7 consulta la tabla, recupera y = 4 y obtiene z = 11. La sentencia y = y + 1 usa el valor previo 4 y lo reemplaza por 5. Cada aparición del identificador conserva en el árbol la anotación de su punto de análisis.')
p('En a = 24.0 + 4 - 1/3*2 + 34 - 1, la subexpresión 1/3 se resuelve primero como división entre enteros y produce 0. El resultado completo es 61.0. Asignar ese resultado a float no cambia retroactivamente los tipos de las subexpresiones.')

page('6 Asignaciones conversiones y valores desconocidos')
h('Verificación de asignaciones',2)
p('visitar_asignacion busca el destino y analiza la expresión. Si uno de ellos presenta error, no agrega un diagnóstico derivado. Si los tipos coinciden, transfiere el valor. Si el destino es float y la expresión es int, convierte el valor conocido mediante float y marca int→float en el nodo de asignación.')
table(['Asignación', 'Comportamiento'],[
('int ← int o float ← float','Válida. Actualiza símbolo, identificador destino y nodo =.'),
('float ← int','Válida con promoción implícita int→float.'),
('int ← float','SEM04, incluso si el valor es 2.0 o es desconocido.'),
('bool ← número o número ← bool','SEM03 por incompatibilidad.'),
('Destino o expresión con error','El padre permanece error sin un diagnóstico adicional.')],[2.5,4.4])
p('Una asignación inválida no sobrescribe el valor anterior de la variable. Por ejemplo, después de y = -2, la instrucción y = 14.54 genera SEM04 y y conserva -2. El nodo = inválido muestra tipo error y valor null.')
h('Representación de null',2)
p('Usamos None de Python para representar un valor que no se conoce en esta fase. El IDE lo convierte en el texto null. No lo confundimos con cero ni con False: una expresión puede tener tipo válido y valor desconocido.')
code('int x, y;\ny = x + 2;\n\nx: tipo int, valor null\nx + 2: tipo int, valor null\ny: tipo int, valor null; sin error por no tener valor')
p('No inicializar una variable no genera un error semántico bajo la regla adoptada para este proyecto. Si una variable float desconocida se asigna a int, la incompatibilidad sigue detectándose porque la validación depende de los tipos, no de conocer el valor.')
p('Si un operando tiene valor None, la operación conserva None. Excepciones aritméticas como división entre cero también dejan el valor desconocido. La versión actual no emite un código específico por división entre cero.')

page('7 Selección ciclos entrada y salida')
h('Selección y unión de estados',2)
p('visitar_if evalúa la condición y guarda el estado inicial de los símbolos. Analiza then en su ámbito, conserva su estado y restaura los valores iniciales antes de revisar else. Ambas ramas se revisan para detectar errores. Si la condición es conocida se conserva la rama seleccionada; si es desconocida solo se conservan los valores iguales entre ambas ramas.')
code('if 2 > 3 then        → condición false\n    y = a + 3;      → se revisa y reporta SEM04\nelse                 → camino seleccionado\n    if 4 > 2 && true then  → true && true = true\n        b = 3.2;     → valor seleccionado para b\n    else\n        b = 5.0;     → también se revisa semánticamente\n    end\nend')
p('El literal true del if interior es independiente del resultado del if exterior. Que una rama aparezca anotada en el árbol no indica que se haya ejecutado. El objetivo del análisis es comprobar todo el código, incluidas las ramas que no se seleccionarían.')
h('Iteración y repetición',2)
p('validar_condicion establece el tipo esperado bool y reporta SEM07 cuando obtiene otro tipo válido. visitar_while revisa primero la condición; visitar_do revisa el cuerpo antes de la condición de do-while o do-until. Los cuerpos se procesan una vez para verificar sus construcciones.')
p('Antes de analizar un ciclo, invalidar_escrituras marca como desconocidos los símbolos que detecta como destinos de asignación, entrada o actualización en el cuerpo. revision ayuda a identificar escrituras al cerrar el ciclo. Si la condición inicial de while es false se recuperan los valores iniciales; en otros casos las escrituras se mantienen conservadoramente como null.')
h('Entrada y salida',2)
p('cin busca el identificador y deja su valor en None porque el dato se recibirá en ejecución. cout analiza las expresiones de salida; las cadenas no requieren una búsqueda en la tabla. Usar mas sin declararlo produce SEM01 tanto en una asignación como en cin.')

page('8 Incrementos decrementos y diagnóstico de errores')
p('El parser mantiene el identificador como hoja y agrega actualizacion con valor +1 o −1 y postfijo para distinguir la posición del operador. El analizador busca el símbolo, comprueba que sea numérico y actualiza su valor conocido. Si es bool, reporta SEM05; si vale null, conserva null.')
p('En una sentencia independiente, a++ muestra el valor final en el árbol: 61.0 pasa a 62.0. En c--, c continúa en null porque no se conoce su valor previo. En una expresión como y = x++, el resultado aportado a y es el valor previo de x, aunque el símbolo x ya queda incrementado. En ++x se aporta el valor actualizado.')
table(['Código', 'Diagnóstico y ejemplo'],[
('SEM01','Identificador no declarado: suma = 45 sin declaración.'),
('SEM02','Duplicado en el mismo ámbito: int x; float x;'),
('SEM03','Asignación incompatible: int x; x = true;'),
('SEM04','Conversión no permitida: int x; x = 32.32;'),
('SEM05','Operador indebido: true + 1, módulo con float o incremento de bool.'),
('SEM06','Igualdad incompatible: 2 == true.'),
('SEM07','Condición no booleana: if 3 then ... end.')],[1,5.9])
h('Registro y prevención de cascadas',2)
p('error añade una tupla con código, línea, columna y descripción a errores, y marca error_sem en el nodo responsable. Cuando falta la ubicación, usa un guion. El tipo interno error permite propagar una falla sin volver a reportarla en todos sus antecesores.')
p('Por ejemplo, en x = z + 2, si z no está declarado, se genera SEM01 en z. La suma y la asignación quedan con tipo error, pero no se añaden falsos errores de operador o de conversión provocados por el mismo identificador. El recorrido continúa con las siguientes sentencias.')

source = (ROOT/'TEST 2026.txt').read_text(encoding='utf-8')
ast, syntax_errors = parse_code(source)
annotated, symbols, errors = analizar_semantica(ast)
page('9 Validación con el programa de prueba')
p('La consigna solicita validar con TestSemantica.txt. En el repositorio, el programa utilizado se llama TEST 2026.txt y su contenido corresponde al código proporcionado en la solicitud. Los números de línea de esta sección consideran main en la línea 1, sin líneas vacías iniciales. Antes de entregar, se debe confirmar que coincide con el archivo descargado del Aula Virtual.')
p('La comprobación local produjo cero errores léxicos, cero errores sintácticos, seis símbolos y seis errores semánticos. El archivo de prueba no se modificó para obtener estos resultados.')
table(['Línea', 'Columna', 'Código', 'Motivo'],[(e[1],e[2],e[0],e[3].split(': ',1)[1]) for e in errors],[.5,.65,.8,4.95])
h('Resultados intermedios representativos',2)
table(['Línea', 'Sentencia o condición', 'Resultado'],[
('6','x = 23','23'),('7','y = 2 + 3 - 1','4'),('8','z = y + 7','11'),
('9','y = y + 1','5'),('10','a = 24.0 + 4 - 1/3*2 + 34 - 1','61.0'),
('11','x = (5-3)*(8/2)','8'),('12','y = 5+3-2*4/7-9','−2'),
('13','z = 8/2 + 15*4','64'),('15','2 > 3','false'),
('18','4 > 2 && true','true'),('25','a++','62.0'),('26','c--','null')],[.5,4.9,1.5])

page('10 Tabla resultante y pruebas automatizadas')
p('Al concluir el recorrido, la tabla conserva los siguientes valores conocidos. x e y terminan como null por el tratamiento conservador de escrituras en los ciclos; esto no afirma que el programa en ejecución les asigne un valor nulo. Las seis variables pertenecen a main y reservan 24 bytes en total.')
table(['Nombre','Tipo','Dirección','Valor final','Líneas de declaración y uso'],[
(s.nombre,s.tipo,s.direccion,'null' if s.valor is None else s.valor,', '.join(map(str,s.lineas))) for s in symbols.entradas],[.6,.65,.8,.9,3.95])
p('b conserva 3.2 por la selección del if interior. Su registro de usos incluye también la asignación de la rama else porque las dos ramas se analizan. a conserva 62.0 después del incremento y c permanece desconocida.')
h('Pruebas reproducidas',2)
p('Se ejecutó python -m unittest test_semantico -v. Las diez pruebas disponibles pasaron en la revisión local. Cada prueba contrasta resultados o diagnósticos concretos, no solo que el analizador termine sin lanzar una excepción.')
table(['Prueba', 'Comportamiento verificado'],[
('Evaluación paso a paso','Valores 5 y 4 en nodos intermedios y z/y dependiente de una asignación.'),
('Conversión y asignación inválida','int→float, SEM04 y conservación del valor anterior.'),
('Valor null y tipo','Operación válida desconocida y rechazo float→int.'),
('Operadores y cascada','SEM05, SEM01 y SEM06 sin reportes derivados.'),
('Entrada y actualizaciones','cin invalida valores conocidos.'),
('Incremento y decremento','La sentencia muestra el valor final en el árbol.'),
('Postfijo en expresión','La expresión utiliza el valor previo y actualiza el símbolo.'),
('Ramas y ciclos','Unión de estados y while inicialmente falso.'),
('División','Truncamiento hacia cero y valor desconocido ante división por cero.'),
('Archivo de prueba','Ausencia de errores previos y los seis diagnósticos esperados.')],[2.1,4.8])

page('11 Integración en el IDE y evidencias del árbol')
p('run_semantic_analysis ejecuta las fases, obtiene la copia anotada y actualiza los paneles mediante build_annotated_tree_view, build_symbol_table_view y build_error_panel. El editor conecta los paneles de errores con los botones de análisis.')
p('El árbol muestra lexema, tipo, dtype, valor, conversión, ámbito, dirección, posición y error. La función de presentación distingue un atributo inexistente de un valor None: el primero queda vacío y el segundo se muestra como null. La tabla presenta los símbolos conservados en entradas.')
capture(1,'Comparación del AST original y el árbol anotado','Insertar una captura del AST original y otra del mismo fragmento en Semántico. Mantener visible y = 2 + 3 - 1 y los valores 5 y 4 de sus operadores. Puede usarse una imagen compuesta con ambas vistas.')
capture(2,'Conversión de tipos y actualización del valor','Mostrar a = 24.0 + 4 - 1/3*2 + 34 - 1 con resultado 61.0 y a++ con resultado 62.0. Mantener visible la columna Conversión en un operando promovido de int a float.')

page('12 Evidencias de símbolos errores y validación')
capture(3,'Tabla de símbolos del programa de prueba','Abrir Hash Table después de analizar el programa. Mostrar los seis identificadores, tipos, direcciones, valores, ámbitos y líneas. Comparar los datos con la sección 10.')
capture(4,'Errores semánticos y ejecución de pruebas','Mostrar los seis diagnósticos de Err. Semánticos con código, línea y columna. Añadir, si se desea, una segunda imagen del terminal con python -m unittest test_semantico -v, el conteo de diez pruebas y el resultado OK.')
p('Las figuras quedan pendientes de insertar. Las tablas de resultados de este documento provienen de la ejecución local del código, pero no sustituyen las capturas solicitadas como evidencia visual del IDE.')

page('13 Alcance limitaciones y preparación de la entrega')
h('Alcance de la implementación',2)
p('El analizador comprueba declaraciones, referencias, asignaciones, operadores y condiciones. Calcula valores cuando cuenta con información suficiente, conserva la trazabilidad de los símbolos y agrega anotaciones sin modificar el AST original. La distinción entre tipo y valor permite seguir comprobando expresiones cuyo resultado todavía es desconocido.')
p('El seguimiento de valores de ciclos es una aproximación conservadora de un recorrido, no una simulación de ejecución ni un análisis iterativo hasta punto fijo. invalidar_escrituras explora destinos de manera estructural; no constituye un análisis completo de flujo ante todas las combinaciones de ámbitos y efectos laterales. Los nodos break se aceptan sin una validación contextual adicional en esta versión.')
p('Las operaciones lógicas revisan sus operandos para detectar errores; el recorrido no implementa un intérprete con ejecución por cortocircuito. La división entre cero deja null sin diagnóstico dedicado. Estas decisiones describen el comportamiento actual y delimitan lo que puede afirmarse a partir del árbol anotado.')
h('Evidencia del proceso de trabajo',2)
p('El diseño de atributos y reglas está registrado en docs/diseno_semantico.md; las notas de implementación y ajustes están en docs/notas_documento.md. El código muestra la separación entre tabla, recorrido y presentación. Las pruebas reproducibles verifican la evaluación, las conversiones, las ramas y la corrección del valor mostrado por incremento y decremento.')
p('No se incluyen capturas inventadas ni se atribuyen a una sesión histórica imágenes generadas después. Para completar la evidencia, se deben insertar las figuras en los espacios anteriores y conservar visibles el código fuente y los paneles correspondientes.')
h('Preparación de la entrega',2)
p('Completar nombres, matrículas, grupo y fecha de la portada. Insertar las capturas y verificar que sus líneas coincidan con el programa entregado. Confirmar el contenido de TestSemantica.txt del Aula Virtual y exportar este Word a PDF.')
p('El archivo comprimido solicitado debe llamarse Fase3_Apellido1_Apellido2_... e incluir el código fuente del analizador semántico con sus dependencias del proyecto, el PDF técnico con evidencias y un archivo de texto con los nombres y matrículas. La consigna compartida establece un descuento de un punto por día de retraso; no proporciona aquí una fecha concreta para completar.')

page('Anexo A Programa utilizado en la validación')
p('Transcripción del archivo local TEST 2026.txt. Se conserva la numeración usada para relacionar cada diagnóstico con su sentencia; los números son solo una ayuda documental y no forman parte del lenguaje.')
lines = source.splitlines()
code('\n'.join(f'{i:02d}  {line}' for i,line in enumerate(lines,1)))

doc.core_properties.title = 'Documentación técnica del análisis semántico del Compilador KC'
doc.core_properties.subject = 'Compiladores II Fase 3'
doc.core_properties.author = 'Equipo del Compilador KC'
doc.core_properties.keywords = 'AST, análisis semántico, tabla de símbolos, tipos, valores'
out = ROOT/'docs'/'Documentacion_Analisis_Semantico_KC.docx'
doc.save(out)
print(out)
